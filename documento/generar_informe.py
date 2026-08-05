from pathlib import Path
import re

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


RAIZ = Path(__file__).resolve().parents[1]
CARPETA_DOCUMENTO = RAIZ / "documento"
CARPETA_DIAGRAMAS = RAIZ / "diagramas"
CARPETA_RECURSOS = CARPETA_DOCUMENTO / "recursos_generados"

ARCHIVO_CONTENIDO = CARPETA_DOCUMENTO / "contenido-documento.md"
ARCHIVO_PORTADA = CARPETA_DOCUMENTO / "datos-portada.md"
ARCHIVO_MATRIZ = RAIZ / "trazabilidad" / "matriz-trazabilidad.md"
ARCHIVO_HISTORIAL = RAIZ / "evidencia" / "historial-git.txt"
ARCHIVO_ESTRUCTURA = RAIZ / "evidencia" / "estructura-repositorio.txt"
ARCHIVO_REFERENCIA = RAIZ / "evidencia" / "referencia-git.md"

SALIDA = CARPETA_DOCUMENTO / "Informe_UML_ASII-08_Esau_Mendoza.docx"


def limpiar_markdown(texto):
    texto = texto.replace("**", "")
    texto = texto.replace("__", "")
    texto = texto.replace("`", "")
    texto = texto.replace("*", "")
    return texto.strip()


def cargar_portada():
    datos = {}

    for linea in ARCHIVO_PORTADA.read_text(encoding="utf-8").splitlines():
        coincidencia = re.match(r"- \*\*(.+?):\*\*\s*(.*)", linea)

        if coincidencia:
            datos[coincidencia.group(1).strip()] = coincidencia.group(2).strip()

    return datos


def configurar_fuente(run, nombre="Times New Roman", tamano=12, negrita=False):
    run.font.name = nombre
    run.font.size = Pt(tamano)
    run.bold = negrita

    run._element.rPr.rFonts.set(qn("w:ascii"), nombre)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), nombre)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), nombre)
def agregar_encabezado(documento, texto, nivel=1):
    estilo = "Heading 1" if nivel == 1 else "Heading 2"
    tamano = 14 if nivel == 1 else 12

    parrafo = documento.add_paragraph(style=estilo)
    parrafo.paragraph_format.keep_with_next = True

    run = parrafo.add_run(texto)
    configurar_fuente(
        run,
        nombre="Times New Roman",
        tamano=tamano,
        negrita=True,
    )
    run.font.color.rgb = RGBColor(0, 0, 0)

    return parrafo


def configurar_estilos(documento):
    normal = documento.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    titulo = documento.styles["Title"]
    titulo.font.name = "Times New Roman"
    titulo.font.size = Pt(16)
    titulo.font.bold = True

    encabezado1 = documento.styles["Heading 1"]
    encabezado1.font.name = "Times New Roman"
    encabezado1.font.size = Pt(14)
    encabezado1.font.bold = True
    encabezado1.font.color.rgb = RGBColor(0, 0, 0)
    encabezado1._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    encabezado1._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    encabezado1._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    encabezado1.paragraph_format.space_before = Pt(0)
    encabezado1.paragraph_format.space_after = Pt(12)
    encabezado1.paragraph_format.keep_with_next = True

    encabezado2 = documento.styles["Heading 2"]
    encabezado2.font.name = "Times New Roman"
    encabezado2.font.size = Pt(12)
    encabezado2.font.bold = True
    encabezado2.font.color.rgb = RGBColor(0, 0, 0)
    encabezado2._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    encabezado2._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    encabezado2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    encabezado2.paragraph_format.keep_with_next = True

    if "Código de evidencia" not in documento.styles:
        codigo = documento.styles.add_style(
            "Código de evidencia",
            WD_STYLE_TYPE.PARAGRAPH,
        )
    else:
        codigo = documento.styles["Código de evidencia"]

    codigo.font.name = "Courier New"
    codigo.font.size = Pt(8)
    codigo.paragraph_format.line_spacing = 1
    codigo.paragraph_format.space_after = Pt(0)


def configurar_pagina(documento):
    seccion = documento.sections[0]
    seccion.page_width = Inches(8.5)
    seccion.page_height = Inches(11)
    seccion.top_margin = Inches(1)
    seccion.bottom_margin = Inches(1)
    seccion.left_margin = Inches(1)
    seccion.right_margin = Inches(1)
    seccion.different_first_page_header_footer = True


def agregar_numero_pagina(parrafo):
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = parrafo.add_run()

    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")

    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = "PAGE"

    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")

    texto = OxmlElement("w:t")
    texto.text = "1"

    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")

    run._r.extend([inicio, instruccion, separador, texto, fin])
    configurar_fuente(run, tamano=10)


def agregar_toc(documento):
    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titulo = titulo.add_run("ÍNDICE")
    configurar_fuente(run_titulo, tamano=14, negrita=True)

    documento.add_paragraph()

    parrafo = documento.add_paragraph()
    run = parrafo.add_run()

    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")

    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = 'TOC \\o "1-2" \\h \\z \\u'

    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")

    texto = OxmlElement("w:t")
    texto.text = "El índice se actualizará automáticamente en Microsoft Word."

    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")

    run._r.extend([inicio, instruccion, separador, texto, fin])


def agregar_portada(documento, datos):
    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.paragraph_format.space_after = Pt(10)

    run = parrafo.add_run(datos.get("Universidad", "").upper())
    configurar_fuente(run, tamano=14, negrita=True)

    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = parrafo.add_run(datos.get("Facultad", ""))
    configurar_fuente(run, tamano=12)

    for _ in range(3):
        documento.add_paragraph()

    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = parrafo.add_run(datos.get("Curso", "").upper())
    configurar_fuente(run, tamano=13, negrita=True)

    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = parrafo.add_run(datos.get("Docente", ""))
    configurar_fuente(run, tamano=12)

    for _ in range(3):
        documento.add_paragraph()

    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = parrafo.add_run(datos.get("Título", "").upper())
    configurar_fuente(run, tamano=15, negrita=True)

    for _ in range(3):
        documento.add_paragraph()

    campos = [
        ("Nombre", datos.get("Estudiante", "")),
        ("No. de Carné", datos.get("Carné", "")),
        ("Correo", datos.get("Correo", "")),
        ("Repositorio", datos.get("Repositorio", "")),
        ("Rama", datos.get("Rama", "")),
        ("Etiqueta", datos.get("Etiqueta de entrega", "")),
    ]

    for etiqueta, valor in campos:
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        parrafo.paragraph_format.left_indent = Inches(0.55)
        parrafo.paragraph_format.first_line_indent = Inches(0)

        run = parrafo.add_run(f"{etiqueta}: ")
        configurar_fuente(run, tamano=11, negrita=True)

        run = parrafo.add_run(valor)
        configurar_fuente(run, tamano=11)

    documento.add_paragraph()

    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = parrafo.add_run(
        f"{datos.get('Fecha de entrega', '')}\n"
        f"{datos.get('Lugar', '')}"
    )
    configurar_fuente(run, tamano=11)

    documento.add_page_break()


def dividir_imagen(ruta, partes, prefijo):
    CARPETA_RECURSOS.mkdir(parents=True, exist_ok=True)

    imagen = Image.open(ruta)
    ancho, alto = imagen.size
    alto_base = alto // partes
    salidas = []

    for indice in range(partes):
        superior = indice * alto_base
        inferior = alto if indice == partes - 1 else (indice + 1) * alto_base

        recorte = imagen.crop((0, superior, ancho, inferior))
        salida = CARPETA_RECURSOS / f"{prefijo}-parte-{indice + 1}.png"
        recorte.save(salida, "PNG")
        salidas.append(salida)

    imagen.close()
    return salidas


def agregar_figura(documento, ruta, titulo, ancho=6.2):
    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = parrafo.add_run()
    run.add_picture(str(ruta), width=Inches(ancho))

    leyenda = documento.add_paragraph()
    leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    leyenda.paragraph_format.space_after = Pt(8)

    run = leyenda.add_run(titulo)
    configurar_fuente(run, tamano=10)
    run.italic = True


def agregar_tabla_matriz(documento):
    if not ARCHIVO_MATRIZ.exists():
        return

    filas = []

    for linea in ARCHIVO_MATRIZ.read_text(encoding="utf-8").splitlines():
        linea = linea.strip()

        if not linea.startswith("|"):
            continue

        celdas = [limpiar_markdown(celda) for celda in linea.strip("|").split("|")]

        if all(re.fullmatch(r":?-+:?", celda.replace(" ", "")) for celda in celdas):
            continue

        filas.append(celdas)

    if len(filas) < 2:
        return

    agregar_encabezado(documento, "Matriz detallada de trazabilidad", nivel=2)

    columnas = max(len(fila) for fila in filas)
    tabla = documento.add_table(rows=1, cols=columnas)
    tabla.style = "Table Grid"
    tabla.autofit = True

    encabezado = tabla.rows[0].cells

    for indice, valor in enumerate(filas[0]):
        encabezado[indice].text = valor
        encabezado[indice].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for run in encabezado[indice].paragraphs[0].runs:
            configurar_fuente(run, tamano=8, negrita=True)

    encabezado_tr = tabla.rows[0]._tr
    propiedades = encabezado_tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    propiedades.append(repetir)

    for fila in filas[1:]:
        celdas = tabla.add_row().cells

        for indice in range(columnas):
            valor = fila[indice] if indice < len(fila) else ""
            celdas[indice].text = valor
            celdas[indice].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for parrafo in celdas[indice].paragraphs:
                parrafo.paragraph_format.space_after = Pt(0)
                parrafo.paragraph_format.line_spacing = 1

                for run in parrafo.runs:
                    configurar_fuente(run, tamano=7.5)


def agregar_evidencia_textual(documento):
    archivos = [
        ("Historial resumido de Git", ARCHIVO_HISTORIAL),
        ("Estructura del repositorio", ARCHIVO_ESTRUCTURA),
        ("Referencia de Git y GitHub", ARCHIVO_REFERENCIA),
    ]

    for titulo, ruta in archivos:
        if not ruta.exists():
            continue

        agregar_encabezado(documento, titulo, nivel=2)

        contenido = ruta.read_text(encoding="utf-8", errors="replace").strip()

        for linea in contenido.splitlines():
            parrafo = documento.add_paragraph(style="Código de evidencia")
            parrafo.add_run(linea)


def extraer_secciones():
    texto = ARCHIVO_CONTENIDO.read_text(encoding="utf-8")
    lineas = texto.splitlines()

    secciones = []
    actual = None

    for linea in lineas:
        if re.match(r"^##\s+\d+\.", linea):
            if actual:
                secciones.append(actual)

            actual = {
                "titulo": limpiar_markdown(linea.replace("##", "", 1)),
                "lineas": [],
            }
            continue

        if actual is not None:
            actual["lineas"].append(linea)

    if actual:
        secciones.append(actual)

    return secciones


def agregar_contenido_seccion(documento, seccion):
    titulo = seccion["titulo"]
    numero = titulo.split(".", 1)[0].strip()

    agregar_encabezado(documento, titulo, nivel=1)

    for linea in seccion["lineas"]:
        linea = linea.rstrip()

        if not linea.strip():
            continue

        if linea.startswith("### "):
            agregar_encabezado(documento, limpiar_markdown(linea[4:]), nivel=2)
            continue

        if re.match(r"^\s*[-*]\s+", linea):
            texto = re.sub(r"^\s*[-*]\s+", "", linea)
            parrafo = documento.add_paragraph(
                limpiar_markdown(texto),
                style="List Bullet",
            )
            parrafo.paragraph_format.line_spacing = 1.5
            continue

        if re.match(r"^\s*\d+\.\s+", linea):
            texto = re.sub(r"^\s*\d+\.\s+", "", linea)
            parrafo = documento.add_paragraph(
                limpiar_markdown(texto),
                style="List Number",
            )
            parrafo.paragraph_format.line_spacing = 1.5
            continue

        if linea.startswith("|"):
            continue

        parrafo = documento.add_paragraph(limpiar_markdown(linea))
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parrafo.paragraph_format.first_line_indent = Inches(0.5)

    if numero == "4":
        documento.add_page_break()
        agregar_figura(
            documento,
            CARPETA_DIAGRAMAS / "casos-de-uso.png",
            "Figura 1. Diagrama de casos de uso del módulo ASII-08.",
        )

    elif numero == "5":
        partes = dividir_imagen(
            CARPETA_DIAGRAMAS / "actividad.png",
            2,
            "actividad",
        )

        for indice, ruta in enumerate(partes, start=1):
            documento.add_page_break()
            agregar_figura(
                documento,
                ruta,
                f"Figura 2.{indice}. Diagrama de actividad, parte {indice} de {len(partes)}.",
            )

    elif numero == "6":
        partes = dividir_imagen(
            CARPETA_DIAGRAMAS / "secuencia.png",
            4,
            "secuencia",
        )

        for indice, ruta in enumerate(partes, start=1):
            documento.add_page_break()
            agregar_figura(
                documento,
                ruta,
                f"Figura 3.{indice}. Diagrama de secuencia, parte {indice} de {len(partes)}.",
            )

    elif numero == "7":
        documento.add_page_break()
        agregar_tabla_matriz(documento)

    elif numero == "9":
        documento.add_page_break()
        agregar_evidencia_textual(documento)


def crear_documento():
    datos = cargar_portada()
    documento = Document()

    configurar_pagina(documento)
    configurar_estilos(documento)

    propiedades = documento.core_properties
    propiedades.title = datos.get("Título", "")
    propiedades.subject = "Diagramas UML del módulo ASII-08"
    propiedades.author = datos.get("Estudiante", "")
    propiedades.keywords = "UML, ASII-08, traslados, altas, camas, PlantUML"

    agregar_portada(documento, datos)
    agregar_toc(documento)
    documento.add_page_break()

    secciones = extraer_secciones()

    for indice, seccion in enumerate(secciones):
        if indice > 0:
            documento.add_page_break()

        agregar_contenido_seccion(documento, seccion)

    for seccion in documento.sections:
        seccion.page_width = Inches(8.5)
        seccion.page_height = Inches(11)
        seccion.top_margin = Inches(1)
        seccion.bottom_margin = Inches(1)
        seccion.left_margin = Inches(1)
        seccion.right_margin = Inches(1)

        agregar_numero_pagina(seccion.footer.paragraphs[0])

    documento.save(SALIDA)
    print(f"Documento generado: {SALIDA}")
    print(f"Tamaño: {SALIDA.stat().st_size} bytes")


if __name__ == "__main__":
    crear_documento()
